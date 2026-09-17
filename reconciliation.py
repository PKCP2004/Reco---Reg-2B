from __future__ import annotations

from io import BytesIO
from pathlib import Path
import re
from datetime import date
from typing import Iterable
from difflib import SequenceMatcher

import pandas as pd
from openpyxl.utils import get_column_letter


# ---------------------------------------------------------------------------
# Column recognition
# ---------------------------------------------------------------------------
# NOTE: "gstin" is deliberately resolved BEFORE "supplier_name" so that a
# header such as "Supplier GSTIN" is claimed by the GSTIN field and can never
# be swallowed by the looser "supplier" alias of the supplier-name field.
FIELD_ALIASES = {
    "gstin": ["gstin", "gst no", "gst number", "supplier gstin", "gstin of supplier", "ctin"],
    "supplier_name": [
        "supplier name", "name of supplier", "name of the supplier", "supplier trade name",
        "trade name", "legal name", "party name", "vendor name", "seller name",
        "trdnm", "lgnm", "supplier", "vendor", "party", "seller",
    ],
    "invoice_no": ["invoice no", "invoice number", "inv no", "document no", "bill no", "inum"],
    "invoice_date": ["invoice date", "date", "document date", "bill date", "idt"],
    "two_b_month": ["2b month", "return period", "period", "month", "tax period"],
    "taxable": ["taxable value", "taxable", "taxable amount"],
    "igst": ["igst", "igst amount"],
    "cgst": ["cgst", "cgst amount"],
    "sgst": ["sgst", "sgst amount", "utgst", "sgst/utgst"],
    "cess": ["cess", "cess amount"],
    "total": ["invoice value", "invoice amount", "total value", "total amount"],
}

VALUE_FIELDS = ["taxable", "igst", "cgst", "sgst", "cess", "total"]
COMPARISON_FIELDS = ["gstin", "supplier_name", "invoice_no", "invoice_date", "two_b_month", *VALUE_FIELDS]

# Statuses that are neither a clean match nor a one-sided item.
REVIEW_STATUSES = [
    "Value Mismatch",
    "GSTIN Mismatch",
    "Same PAN - Different GSTIN",
    "Tax Head Mismatch (IGST vs CGST/SGST)",
    "Likely Match - Register Invoice Error",
    "Likely Match - Invoice Format",
]

STANDARD_GST_RATES = [0.0, 0.1, 0.25, 1.0, 1.5, 3.0, 5.0, 6.0, 7.5, 12.0, 18.0, 28.0]
_GST_CODE = "0123456789ABCDEFGHIJKLMNOPQRSTUVWXYZ"
_GSTIN_PATTERN = re.compile(r"^[0-9]{2}[A-Z]{5}[0-9]{4}[A-Z][0-9A-Z]Z[0-9A-Z]$")
_NAME_NOISE = re.compile(
    r"\b(PRIVATE|PVT|LIMITED|LTD|LLP|COMPANY|CORPORATION|CORP|INC|AND|THE|MS)\b"
)


# ---------------------------------------------------------------------------
# Cleaning helpers
# ---------------------------------------------------------------------------
def _clean_text(value) -> str:
    if pd.isna(value):
        return ""
    return re.sub(r"[^A-Z0-9]", "", str(value).upper())


def _display_text(value) -> str:
    if pd.isna(value):
        return ""
    return re.sub(r"\s+", " ", str(value)).strip()


def _name_key(value) -> str:
    """Normalize a supplier name for comparison.

    Uppercases, drops punctuation and the usual corporate suffixes (PVT, LTD,
    LLP, ...) so "ABC Traders Pvt. Ltd." and "ABC TRADERS PRIVATE LIMITED"
    compare as the same supplier.
    """
    text = _display_text(value).upper()
    if not text:
        return ""
    text = re.sub(r"[^A-Z0-9 ]", " ", text)
    text = _NAME_NOISE.sub(" ", text)
    return re.sub(r"\s+", "", text)


def _number(value) -> float:
    """Parse a numeric cell into a rounded float.

    Handles thousands separators, currency symbols/text, and accounting-style
    negative numbers written in parentheses, e.g. "(1,234.50)" -> -1234.50.
    """
    if pd.isna(value) or str(value).strip() == "":
        return 0.0
    text = str(value).strip()
    is_parenthesised_negative = text.startswith("(") and text.endswith(")")
    if is_parenthesised_negative:
        text = text[1:-1]
    text = re.sub(r"(?i)\b(rs|inr|rupees?)\b\.?", "", text)
    text = text.replace("\u20b9", "").replace(",", "")
    match = re.search(r"-?\d+(?:\.\d+)?", text)
    if not match:
        return 0.0
    try:
        result = float(match.group())
    except ValueError:
        return 0.0
    if is_parenthesised_negative and result > 0:
        result = -result
    return round(result, 2)


def _pan(gstin: str) -> str:
    """The 10-character PAN embedded in a 15-character GSTIN."""
    return gstin[2:12] if len(gstin) == 15 else ""


def _invoice_core(invoice_raw) -> str:
    """A format-insensitive core of an invoice number.

    "INV-1001", "INV/1001/25-26", "0001001" and "1001" all reduce to the same
    core, which matches the very common case where the register stores a short
    invoice number while GSTR-2B carries the supplier's full series.

    This deliberately works on the ORIGINAL text, before separators are
    stripped, so "INV/1006/26-27" yields the serial 1006 rather than the
    run-together 10062627.
    """
    text = _display_text(invoice_raw)
    if not text:
        return ""
    digit_runs = re.findall(r"\d+", text)
    if not digit_runs:
        return ""
    core = max(digit_runs, key=len).lstrip("0")
    return core if len(core) >= 3 else ""


def _gstin_is_valid(gstin: str) -> bool:
    if not _GSTIN_PATTERN.match(gstin or ""):
        return False
    total = 0
    for index, character in enumerate(gstin[:14]):
        value = _GST_CODE.index(character)
        product = value * ((index % 2) + 1)
        total += product // 36 + product % 36
    return _GST_CODE[(36 - total % 36) % 36] == gstin[14]


def _find_column(columns: Iterable, aliases: list[str], exclude: set | None = None):
    """Find the best-matching source column for a field's alias list.

    Exact (case-insensitive) header matches always beat partial matches, and
    columns already claimed by an earlier field are skipped so one source
    column can never be mapped into two standardized fields.
    """
    exclude = exclude or set()
    normalized = {
        str(column).strip().lower(): column
        for column in columns
        if column not in exclude
    }
    for alias in aliases:
        if alias in normalized:
            return normalized[alias]
    for alias in aliases:
        for name, original in normalized.items():
            if alias in name:
                return original
    return None


def standardize_frame(frame: pd.DataFrame, source: str) -> pd.DataFrame:
    frame = frame.copy()
    output = pd.DataFrame(index=frame.index)
    used_columns: set = set()
    for field, aliases in FIELD_ALIASES.items():
        column = _find_column(frame.columns, aliases, exclude=used_columns)
        if column is not None:
            used_columns.add(column)
        output[field] = frame[column] if column is not None else ""

    output["invoice_raw"] = output["invoice_no"].map(_display_text)
    output["gstin"] = output["gstin"].map(_clean_text)
    output["supplier_name"] = output["supplier_name"].map(_display_text)
    output["invoice_no"] = output["invoice_no"].map(_clean_text)
    output["invoice_date"] = (
        pd.to_datetime(output["invoice_date"], errors="coerce", dayfirst=True)
        .dt.strftime("%Y-%m-%d")
        .fillna("")
    )
    output["two_b_month"] = output["two_b_month"].fillna("").astype(str).str.strip()
    for field in VALUE_FIELDS:
        output[field] = output[field].map(_number)

    output["supplier_key"] = output["supplier_name"].map(_name_key)
    output["pan"] = output["gstin"].map(_pan)
    output["invoice_core"] = output["invoice_raw"].map(_invoice_core)
    output["source"] = source
    output["source_row"] = frame.index + 2
    output = output[(output["gstin"] != "") | (output["invoice_no"] != "")].reset_index(drop=True)
    output["match_key"] = output["gstin"] + "|" + output["invoice_no"]
    return output


# ---------------------------------------------------------------------------
# Data-quality / compliance checks
# ---------------------------------------------------------------------------
def _itc_deadline(invoice_date: date) -> date:
    """30 November following the end of the invoice's financial year."""
    fy_start_year = invoice_date.year if invoice_date.month >= 4 else invoice_date.year - 1
    return date(fy_start_year + 1, 11, 30)


QUALITY_COLUMNS = [
    "Source", "Source Row", "GSTIN", "Supplier Name", "Invoice No", "Severity", "Check", "Detail",
]


def run_quality_checks(frame: pd.DataFrame, label: str, tolerance: float = 1.0) -> pd.DataFrame:
    """Row-level sanity checks that are independent of the matching outcome."""
    today = date.today()
    issues: list[dict] = []
    if frame is None or frame.empty:
        return pd.DataFrame(columns=QUALITY_COLUMNS)

    duplicate_keys = set(frame[frame.duplicated("match_key", keep=False)]["match_key"])

    for row in frame.to_dict("records"):
        def flag(severity: str, check: str, detail: str, row=row):
            issues.append({
                "Source": label,
                "Source Row": row["source_row"],
                "GSTIN": row["gstin"],
                "Supplier Name": row["supplier_name"],
                "Invoice No": row["invoice_no"],
                "Severity": severity,
                "Check": check,
                "Detail": detail,
            })

        if not row["gstin"]:
            flag("High", "GSTIN missing", "No GSTIN on this line.")
        elif not _gstin_is_valid(row["gstin"]):
            flag("High", "GSTIN invalid", "Fails GSTIN format / check-digit validation.")

        if not row["supplier_name"]:
            flag("Low", "Supplier name missing", "Add the supplier name for easier review.")

        if not row["invoice_no"]:
            flag("High", "Invoice number missing", "No invoice number on this line.")

        if row["match_key"] in duplicate_keys:
            twins = frame[(frame["match_key"] == row["match_key"]) & (frame["source_row"] != row["source_row"])]
            identical = any(
                all(abs(row[field] - twin[field]) <= tolerance for field in VALUE_FIELDS)
                for twin in twins.to_dict("records")
            )
            if identical:
                flag("High", "Duplicate invoice in file",
                     "Same GSTIN + invoice number with the same values - double ITC claim risk.")
            else:
                flag("Medium", "Repeated invoice number",
                     "Same GSTIN + invoice number on several lines - confirm it is a rate-wise split, "
                     "not a duplicate booking.")

        taxes = row["igst"] + row["cgst"] + row["sgst"]
        if row["igst"] > tolerance and (row["cgst"] > tolerance or row["sgst"] > tolerance):
            flag("High", "IGST with CGST/SGST",
                 "IGST and CGST/SGST cannot both apply to the same invoice line.")
        if abs(row["cgst"] - row["sgst"]) > tolerance:
            flag("High", "CGST not equal to SGST",
                 f"CGST {row['cgst']:.2f} vs SGST {row['sgst']:.2f}.")

        if row["total"] and abs(row["taxable"] + taxes + row["cess"] - row["total"]) > max(tolerance, 1.0):
            flag("Medium", "Invoice value does not add up",
                 f"Taxable + taxes = {row['taxable'] + taxes + row['cess']:.2f}, "
                 f"invoice value = {row['total']:.2f}.")

        if row["taxable"] > 0 and taxes >= 0:
            implied = round(taxes / row["taxable"] * 100, 2)
            if not any(abs(implied - rate) <= 0.2 for rate in STANDARD_GST_RATES):
                flag("Medium", "Non-standard tax rate",
                     f"Implied rate {implied:.2f}% is not a notified GST rate.")

        if not row["invoice_date"]:
            flag("Medium", "Invoice date missing or unreadable", "Date could not be parsed.")
        else:
            invoice_date = date.fromisoformat(row["invoice_date"])
            if invoice_date > today:
                flag("High", "Future invoice date", f"Invoice dated {row['invoice_date']}.")
            elif today > _itc_deadline(invoice_date):
                flag("High", "ITC time limit likely expired",
                     f"Section 16(4) cut-off was {_itc_deadline(invoice_date).isoformat()}.")

        if min(row["taxable"], row["total"]) < 0:
            flag("Low", "Negative values", "Looks like a credit note / debit note line.")

    return pd.DataFrame(issues, columns=QUALITY_COLUMNS)


# ---------------------------------------------------------------------------
# Reconciliation
# ---------------------------------------------------------------------------
def reconcile(gstr2b: pd.DataFrame, register: pd.DataFrame, tolerance: float = 1.0,
              progress_callback=None) -> dict[str, pd.DataFrame]:
    two_b = gstr2b.copy().reset_index(drop=True)
    inward = register.copy().reset_index(drop=True)
    two_b["_side"] = "GSTR-2B"
    inward["_side"] = "Inward Register"
    two_b["_row_id"] = two_b.index
    inward["_row_id"] = inward.index

    # --- Supplier directory: GSTIN -> supplier name, learned from GSTR-2B ---
    supplier_directory: dict[str, str] = {}
    for frame in (two_b, inward):
        named = frame[(frame["supplier_name"] != "") & (frame["gstin"] != "")]
        for gstin, group in named.groupby("gstin"):
            supplier_directory.setdefault(gstin, group["supplier_name"].mode().iloc[0])

    for frame in (two_b, inward):
        frame["supplier_name_source"] = frame["supplier_name"].map(lambda name: "File" if name else "")
        blank = frame["supplier_name"] == ""
        if blank.any():
            frame.loc[blank, "supplier_name"] = frame.loc[blank, "gstin"].map(
                lambda gstin: supplier_directory.get(gstin, "")
            )
            frame.loc[blank & (frame["supplier_name"] != ""), "supplier_name_source"] = "Derived from GSTIN"
        frame.loc[frame["supplier_name"] == "", "supplier_name_source"] = "Unknown"
        frame["supplier_key"] = frame["supplier_name"].map(_name_key)

    for column in ["reco_status", "likely_2b_invoice_no", "matched_2b_month",
                   "matched_2b_link", "reco_basis", "matched_2b_supplier_name"]:
        inward[column] = ""
    inward["matched_2b_row_id"] = pd.Series([None] * len(inward), index=inward.index, dtype="object")
    inward["reco_confidence"] = pd.Series([None] * len(inward), index=inward.index, dtype="object")

    matched_register_rows: set = set()
    matched_two_b_rows: set = set()
    results: list[dict] = []

    two_b_records = {row["_row_id"]: row for row in two_b.to_dict("records")}
    inward_records = inward.to_dict("index")

    # ------------------------------------------------------------------
    # Shared helpers
    # ------------------------------------------------------------------
    def comparison_fields(two_b_row=None, register_row=None):
        values = {}
        for field in COMPARISON_FIELDS:
            values[f"2B_{field}"] = "" if two_b_row is None else two_b_row.get(field, "")
            values[f"Register_{field}"] = "" if register_row is None else register_row.get(field, "")
        return values

    def value_differences(two_b_row, register_row):
        return {
            field: round(two_b_row[field] - register_row[field], 2)
            for field in VALUE_FIELDS
            if abs(two_b_row[field] - register_row[field]) > tolerance
        }

    def date_gap(two_b_row, register_row):
        left, right = two_b_row.get("invoice_date", ""), register_row.get("invoice_date", "")
        if not left or not right:
            return None
        return abs((date.fromisoformat(left) - date.fromisoformat(right)).days)

    def date_warning(two_b_row, register_row) -> str:
        if two_b_row is None or register_row is None:
            return ""
        left, right = two_b_row.get("invoice_date", ""), register_row.get("invoice_date", "")
        if not left or not right or left == right:
            return ""
        return f"Date differs by {date_gap(two_b_row, register_row)} day(s): 2B {left} | Register {right}"

    def supplier_warning(two_b_row, register_row) -> str:
        if two_b_row is None or register_row is None:
            return ""
        left, right = two_b_row.get("supplier_key", ""), register_row.get("supplier_key", "")
        if not left or not right or left == right:
            return ""
        if SequenceMatcher(None, left, right).ratio() >= 0.85:
            return ""
        return (f"Supplier name differs: 2B '{two_b_row.get('supplier_name', '')}' | "
                f"Register '{register_row.get('supplier_name', '')}'")

    def tax_head_swapped(two_b_row, register_row) -> bool:
        """IGST on one side but CGST+SGST of the same amount on the other."""
        if abs(two_b_row["taxable"] - register_row["taxable"]) > tolerance:
            return False
        for left, right in ((two_b_row, register_row), (register_row, two_b_row)):
            if left["igst"] > tolerance and right["igst"] <= tolerance:
                if abs(left["igst"] - (right["cgst"] + right["sgst"])) <= max(tolerance, 2.0):
                    return True
        return False

    def append_result(two_b_row=None, register_row=None, status="", differences=None,
                      confidence=0, basis="", action=""):
        primary = two_b_row if two_b_row is not None else register_row
        differences = differences or {}
        result = {
            **dict(primary),
            **comparison_fields(two_b_row, register_row),
            "status": status,
            "confidence": confidence,
            "match_basis": basis,
            "difference": "; ".join(f"{field}: {value:+.2f}" for field, value in differences.items()),
            "value_difference_total": round(sum(abs(value) for value in differences.values()), 2),
            "date_warning": date_warning(two_b_row, register_row),
            "supplier_name_warning": supplier_warning(two_b_row, register_row),
            "register_row": "" if register_row is None else str(register_row["source_row"]),
            "two_b_row": "" if two_b_row is None else str(two_b_row["source_row"]),
            "recommended_action": action,
        }
        results.append(result)

    def annotate_register(register_row, two_b_row, status, confidence, basis):
        row_id = register_row["_row_id"]
        updates = {
            "reco_status": status,
            "likely_2b_invoice_no": two_b_row["invoice_no"] if two_b_row is not None else "",
            "matched_2b_month": two_b_row["two_b_month"] if two_b_row is not None else "",
            "matched_2b_supplier_name": two_b_row["supplier_name"] if two_b_row is not None else "",
            "matched_2b_row_id": two_b_row["_row_id"] if two_b_row is not None else None,
            "matched_2b_link": "Open GSTR-2B invoice" if two_b_row is not None else "",
            "reco_confidence": confidence,
            "reco_basis": basis,
        }
        for column, value in updates.items():
            inward.at[row_id, column] = value
            register_row[column] = value

    def combine(records: list[dict]) -> dict:
        combined = dict(records[0])
        combined["source_row"] = ", ".join(str(record["source_row"]) for record in records)
        combined["split_count"] = len(records)
        combined["_row_ids"] = [record["_row_id"] for record in records]
        for field in VALUE_FIELDS:
            combined[field] = round(sum(record[field] for record in records), 2)
        return combined

    def classify(source_row, target, differences, base_basis):
        """Decide status / basis / action for a matched pair."""
        same_gstin = source_row["gstin"] == target["gstin"]
        same_pan = bool(source_row["pan"]) and source_row["pan"] == target["pan"]
        same_invoice = bool(source_row["invoice_no"]) and source_row["invoice_no"] == target["invoice_no"]

        if differences and tax_head_swapped(source_row, target):
            return ("Tax Head Mismatch (IGST vs CGST/SGST)",
                    "Same taxable value; tax charged under different heads",
                    "Confirm place of supply - inter-state vs intra-state treatment")
        if not same_gstin and same_pan:
            return ("Same PAN - Different GSTIN",
                    f"{base_basis}; same PAN, different state registration",
                    "Confirm which GST registration this purchase belongs to")
        if not same_gstin:
            return ("GSTIN Mismatch", f"{base_basis}; GSTIN differs",
                    "Verify the supplier GSTIN recorded in the register")
        if differences:
            return ("Value Mismatch", base_basis, "Review taxable value and tax amounts")
        if not same_invoice:
            same_core = bool(source_row["invoice_core"]) and source_row["invoice_core"] == target["invoice_core"]
            if same_core:
                return ("Likely Match - Invoice Format",
                        f"{base_basis}; same serial, different invoice-number format",
                        f"Align the register invoice number with the 2B series "
                        f"({source_row.get('invoice_raw') or source_row['invoice_no']})")
            return ("Likely Match - Register Invoice Error",
                    f"{base_basis}; invoice number differs",
                    f"Confirm inward invoice number should be "
                    f"{source_row.get('invoice_raw') or source_row['invoice_no']}")
        if date_warning(source_row, target):
            return ("Matched", base_basis, "Review invoice date difference")
        if supplier_warning(source_row, target):
            return ("Matched", base_basis, "Verify supplier name")
        return ("Matched", base_basis, "No action")

    def record_pair(source_row, target, base_basis, confidence):
        differences = value_differences(source_row, target)
        status, basis, action = classify(source_row, target, differences, base_basis)
        for target_id in target.get("_row_ids", [target["_row_id"]]):
            annotate_register(inward_records[target_id], source_row, status, confidence, basis)
        append_result(two_b_row=source_row, register_row=target, status=status,
                      differences=differences, confidence=confidence, basis=basis, action=action)

    def claim(source_row, target):
        for row_id in target.get("_row_ids", [target["_row_id"]]):
            matched_register_rows.add(row_id)
        for row_id in source_row.get("_row_ids", [source_row["_row_id"]]):
            matched_two_b_rows.add(row_id)

    # ------------------------------------------------------------------
    # Phase 1 - deterministic key tiers, strongest first. Each tier is fully
    # resolved across the whole file before the next one starts, so a strong
    # match is never stolen by a weaker one that merely appeared earlier.
    # ------------------------------------------------------------------
    def tier_key(row, tier: str) -> str:
        if tier == "gstin_invoice" and row["gstin"] and row["invoice_no"]:
            return f"{row['gstin']}|{row['invoice_no']}"
        if tier == "gstin_core" and row["gstin"] and row["invoice_core"]:
            return f"{row['gstin']}|#{row['invoice_core']}"
        if tier == "pan_invoice" and row["pan"] and row["invoice_no"]:
            return f"{row['pan']}|{row['invoice_no']}"
        if tier == "pan_core" and row["pan"] and row["invoice_core"]:
            return f"{row['pan']}|#{row['invoice_core']}"
        return ""

    TIERS = [
        ("gstin_invoice", "Exact GSTIN + invoice number", 100),
        ("gstin_core", "Same GSTIN + same invoice serial (different format)", 96),
        ("pan_invoice", "Same PAN + exact invoice number", 90),
        ("pan_core", "Same PAN + same invoice serial (different format)", 86),
    ]

    def pair_group(two_b_ids: list[int], reg_ids: list[int], basis: str, confidence: int):
        two_b_rows = [two_b_records[row_id] for row_id in two_b_ids]
        reg_rows = [inward_records[row_id] for row_id in reg_ids]

        # 1 x N -> either a single clean line, or a split inward invoice.
        if len(two_b_rows) == 1 and len(reg_rows) > 1:
            source_row = two_b_rows[0]
            exact = [row for row in reg_rows if not value_differences(source_row, row)]
            if len(exact) == 1:
                claim(source_row, exact[0])
                record_pair(source_row, exact[0], basis, confidence)
                return
            target = combine(reg_rows)
            claim(source_row, target)
            record_pair(source_row, target,
                        f"Split inward invoice ({len(reg_rows)} lines) summed to 2B", confidence)
            return

        # N x 1 -> a 2B invoice reported in several lines (rate-wise split).
        if len(reg_rows) == 1 and len(two_b_rows) > 1:
            target = reg_rows[0]
            exact = [row for row in two_b_rows if not value_differences(row, target)]
            if len(exact) == 1:
                claim(exact[0], target)
                record_pair(exact[0], target, basis, confidence)
                return
            source_row = combine(two_b_rows)
            claim(source_row, target)
            record_pair(source_row, target,
                        f"Split 2B invoice ({len(two_b_rows)} lines) summed to register", confidence)
            return

        # N x M -> pair greedily on the closest invoice value.
        pairs = sorted(
            ((abs(source["total"] - target["total"]), source["_row_id"], target["_row_id"])
             for source in two_b_rows for target in reg_rows),
            key=lambda item: item,
        )
        for _, two_b_id, reg_id in pairs:
            if two_b_id in matched_two_b_rows or reg_id in matched_register_rows:
                continue
            source_row, target = two_b_records[two_b_id], inward_records[reg_id]
            claim(source_row, target)
            note = basis if len(two_b_rows) == 1 else f"{basis} (repeated invoice number)"
            record_pair(source_row, target, note, confidence)

    for index, (tier, basis, confidence) in enumerate(TIERS):
        if progress_callback:
            progress_callback(0.28 + index * 0.04, f"Matching: {basis.lower()}")
        two_b_groups: dict[str, list[int]] = {}
        for row_id, row in two_b_records.items():
            if row_id in matched_two_b_rows:
                continue
            key = tier_key(row, tier)
            if key:
                two_b_groups.setdefault(key, []).append(row_id)
        reg_groups: dict[str, list[int]] = {}
        for row_id, row in inward_records.items():
            if row_id in matched_register_rows:
                continue
            key = tier_key(row, tier)
            if key:
                reg_groups.setdefault(key, []).append(row_id)
        for key, candidate_two_b_ids in two_b_groups.items():
            reg_ids = [row_id for row_id in reg_groups.get(key, []) if row_id not in matched_register_rows]
            open_two_b_ids = [row_id for row_id in candidate_two_b_ids if row_id not in matched_two_b_rows]
            if open_two_b_ids and reg_ids:
                pair_group(open_two_b_ids, reg_ids, basis, confidence)

    # ------------------------------------------------------------------
    # Phase 2 - weighted fuzzy scoring across every remaining candidate pair.
    # Every pair is scored first and assignment happens strictly in descending
    # score order, so the best available candidate always wins.
    # ------------------------------------------------------------------
    if progress_callback:
        progress_callback(0.50, "Scoring fuzzy invoice candidates")

    register_index: dict[str, list[int]] = {}

    def add_index(key: str, row_id: int):
        if key and not key.endswith(":"):
            register_index.setdefault(key, []).append(row_id)

    def candidate_keys(row) -> list[str]:
        return [
            f"gstin:{row['gstin']}",
            f"pan:{row['pan']}",
            f"inv:{row['invoice_no']}",
            f"core:{row['invoice_core']}",
            f"prefix:{row['invoice_no'][:4]}",
            f"name:{row['supplier_key'][:12]}",
            f"amt:{row['total']:.2f}",
            f"datetax:{row['invoice_date']}|{row['taxable']:.2f}",
        ]

    for row_id, row in inward_records.items():
        if row_id in matched_register_rows:
            continue
        for key in candidate_keys(row):
            add_index(key, row_id)

    def score_pair(source_row, target):
        same_invoice = bool(source_row["invoice_no"]) and source_row["invoice_no"] == target["invoice_no"]
        same_core = bool(source_row["invoice_core"]) and source_row["invoice_core"] == target["invoice_core"]
        similarity = SequenceMatcher(None, source_row["invoice_no"], target["invoice_no"]).ratio()
        same_gstin = bool(source_row["gstin"]) and source_row["gstin"] == target["gstin"]
        same_pan = bool(source_row["pan"]) and source_row["pan"] == target["pan"]
        name_similarity = (
            SequenceMatcher(None, source_row["supplier_key"], target["supplier_key"]).ratio()
            if source_row["supplier_key"] and target["supplier_key"] else 0.0
        )
        differences = value_differences(source_row, target)
        taxable_ok = abs(source_row["taxable"] - target["taxable"]) <= tolerance
        total_ok = abs(source_row["total"] - target["total"]) <= tolerance
        gap = date_gap(source_row, target)

        score = 0
        reasons: list[str] = []
        if same_invoice:
            score += 45
            reasons.append("invoice number")
        elif same_core:
            score += 38
            reasons.append("invoice serial")
        elif similarity >= 0.80:
            score += int(30 * similarity)
            reasons.append(f"similar invoice number ({similarity:.0%})")

        if same_gstin:
            score += 25
            reasons.append("GSTIN")
        elif same_pan:
            score += 15
            reasons.append("PAN")

        if name_similarity >= 0.85:
            score += 8
            reasons.append("supplier name")

        if not differences:
            score += 20
            reasons.append("all tax values")
        elif taxable_ok:
            score += 12
            reasons.append("taxable value")
            if total_ok:
                score += 4

        if gap == 0:
            score += 10
            reasons.append("invoice date")
        elif gap is not None and gap <= 7:
            score += 5

        invoice_anchor = same_invoice or same_core or similarity >= 0.80
        party_anchor = same_gstin or same_pan or name_similarity >= 0.85
        value_anchor = (not differences) or (taxable_ok and total_ok)
        # At least two independent anchors are required. This is what keeps
        # stronger matching from degrading into loose guessing.
        acceptable = score >= 55 and sum([invoice_anchor, party_anchor, value_anchor]) >= 2
        basis = "Matched on " + " + ".join(reasons) if reasons else "Weak candidate"
        return score, basis, acceptable

    candidates = []
    for two_b_id, source_row in two_b_records.items():
        if two_b_id in matched_two_b_rows:
            continue
        candidate_ids: set = set()
        for key in candidate_keys(source_row):
            candidate_ids.update(register_index.get(key, []))
        for row_id in candidate_ids:
            if row_id in matched_register_rows:
                continue
            score, basis, acceptable = score_pair(source_row, inward_records[row_id])
            if acceptable:
                candidates.append((score, two_b_id, row_id, basis))

    candidates.sort(key=lambda item: (-item[0], item[1], item[2]))
    for score, two_b_id, row_id, basis in candidates:
        if two_b_id in matched_two_b_rows or row_id in matched_register_rows:
            continue
        source_row, target = two_b_records[two_b_id], inward_records[row_id]
        claim(source_row, target)
        record_pair(source_row, target, basis, min(95, score))

    # ------------------------------------------------------------------
    # Phase 3 - whatever remains is genuinely one-sided.
    # ------------------------------------------------------------------
    if progress_callback:
        progress_callback(0.80, "Finalizing 2B-only invoices")

    for two_b_id, source_row in two_b_records.items():
        if two_b_id in matched_two_b_rows:
            continue
        append_result(
            two_b_row=source_row, status="In 2B Only", confidence=0,
            basis="No reliable inward-register candidate",
            action="Book the purchase, or check the invoice number / GSTIN in the register",
        )

    if progress_callback:
        progress_callback(0.88, f"Checking {len(inward):,} inward-register rows for register-only items")

    for row_id, register_row in inward_records.items():
        if row_id in matched_register_rows:
            continue
        annotate_register(register_row, None, "In Register Only", 0, "No reliable 2B candidate")
        append_result(register_row=register_row, status="In Register Only", confidence=0,
                      basis="No reliable 2B candidate",
                      action="Follow up with the supplier - invoice not filed or filed in another period")

    # ------------------------------------------------------------------
    # Result assembly
    # ------------------------------------------------------------------
    if progress_callback:
        progress_callback(0.93, "Building reconciliation result tables")

    result = pd.DataFrame(results).drop(
        columns=["_side", "_row_id", "_row_ids"], errors="ignore"
    )
    if "two_b_month" not in result.columns:
        result["two_b_month"] = ""

    if not result.empty:
        result["_register_order"] = pd.to_numeric(
            result["register_row"].astype(str).str.split(",").str[0], errors="coerce"
        )
        result["_register_only"] = result["_register_order"].isna()
        result = result.sort_values(["_register_only", "_register_order"], kind="stable").drop(
            columns=["_register_order", "_register_only"]
        )
        duplicate_2b = set(two_b[two_b.duplicated("match_key", keep=False)]["match_key"])
        duplicate_register = set(inward[inward.duplicated("match_key", keep=False)]["match_key"])
        result["duplicate_2b_invoice"] = result["match_key"].isin(duplicate_2b).map({True: "Yes", False: "No"})
        result["duplicate_register_invoice"] = result["match_key"].isin(duplicate_register).map({True: "Yes", False: "No"})
    else:
        result["duplicate_2b_invoice"] = "No"
        result["duplicate_register_invoice"] = "No"

    # Supplier name sits right next to the GSTIN so the sheet reads naturally.
    preferred = [
        "supplier_name", "gstin", "invoice_no", "invoice_date", "two_b_month",
        "status", "confidence", "match_basis", "difference", "value_difference_total",
        "date_warning", "supplier_name_warning", "recommended_action",
        "two_b_row", "register_row",
    ]
    ordered = [column for column in preferred if column in result.columns]
    result = result[ordered + [column for column in result.columns if column not in ordered]]

    two_b_source = two_b.drop(columns=["_side", "_row_id"], errors="ignore")
    inward_source = inward.drop(columns=["_side", "_row_id"], errors="ignore")

    matched_details = inward_source[inward_source["reco_status"] == "Matched"].copy()
    if not matched_details.empty:
        two_b_details = two_b_source[[
            "match_key", "gstin", "supplier_name", "invoice_no", "invoice_date", "two_b_month", *VALUE_FIELDS,
        ]].drop_duplicates("match_key").rename(columns={
            column: f"2B_{column}" for column in
            ["gstin", "supplier_name", "invoice_no", "invoice_date", "two_b_month", *VALUE_FIELDS]
        })
        matched_details = matched_details.merge(two_b_details, on="match_key", how="left")
        matched_details = matched_details.rename(columns={
            "supplier_name": "Supplier Name",
            "gstin": "Register GSTIN",
            "invoice_no": "Register Invoice No",
            "source_row": "Register Source Row",
            "reco_status": "Status",
            "likely_2b_invoice_no": "Matched 2B Invoice No",
            "reco_confidence": "Match Confidence",
            "reco_basis": "Match Basis",
        })

    monthly_analysis = inward_source.copy()
    monthly_analysis["matched_2b_month"] = monthly_analysis["matched_2b_month"].replace("", "Not found in 2B")
    monthly_analysis = monthly_analysis.groupby(["matched_2b_month", "reco_status"], dropna=False).agg(
        Invoice_Lines=("invoice_no", "size"),
        Taxable_Value=("taxable", "sum"),
        IGST=("igst", "sum"),
        CGST=("cgst", "sum"),
        SGST=("sgst", "sum"),
        Cess=("cess", "sum"),
        Invoice_Value=("total", "sum"),
    ).reset_index().rename(columns={"matched_2b_month": "2B Month", "reco_status": "Status"})

    # --- Supplier-wise summary: the reference view the new column enables ---
    supplier_summary = pd.DataFrame()
    if not result.empty:
        working = result.copy()
        working["Supplier Name"] = working["supplier_name"].replace("", "Unknown supplier")
        working["_matched"] = (working["status"] == "Matched").astype(int)
        working["_review"] = working["status"].isin(REVIEW_STATUSES).astype(int)
        working["_2b_only"] = (working["status"] == "In 2B Only").astype(int)
        working["_reg_only"] = (working["status"] == "In Register Only").astype(int)
        supplier_summary = working.groupby(["Supplier Name", "gstin"], dropna=False).agg(
            Invoices=("status", "size"),
            Matched=("_matched", "sum"),
            Needs_Review=("_review", "sum"),
            In_2B_Only=("_2b_only", "sum"),
            In_Register_Only=("_reg_only", "sum"),
            Taxable_Value=("taxable", "sum"),
            IGST=("igst", "sum"),
            CGST=("cgst", "sum"),
            SGST=("sgst", "sum"),
            Invoice_Value=("total", "sum"),
            Value_Difference=("value_difference_total", "sum"),
        ).reset_index().rename(columns={"gstin": "GSTIN"})
        supplier_summary["Match_Rate_%"] = (
            supplier_summary["Matched"] / supplier_summary["Invoices"] * 100
        ).round(1)
        supplier_summary = supplier_summary.sort_values(
            ["Needs_Review", "In_Register_Only", "Invoice_Value"], ascending=False
        ).reset_index(drop=True)

    quality = pd.concat(
        [
            run_quality_checks(two_b_source, "GSTR-2B", tolerance),
            run_quality_checks(inward_source, "Inward Register", tolerance),
        ],
        ignore_index=True,
    )
    if not quality.empty:
        quality["_order"] = quality["Severity"].map({"High": 0, "Medium": 1, "Low": 2}).fillna(3)
        quality = quality.sort_values(["_order", "Source", "Source Row"]).drop(
            columns=["_order"]
        ).reset_index(drop=True)

    supplier_dir_frame = pd.DataFrame(
        sorted(supplier_directory.items()), columns=["GSTIN", "Supplier Name"]
    )

    empty = result.iloc[0:0]
    return {
        "reconciliation": result,
        "gstr2b_source": two_b_source,
        "inward_register_source": inward_source,
        "matched_details": matched_details,
        "monthly_analysis": monthly_analysis,
        "supplier_summary": supplier_summary,
        "quality_checks": quality,
        "supplier_directory": supplier_dir_frame,
        "matched": result[result["status"] == "Matched"].copy() if not result.empty else empty,
        "mismatches": result[result["status"].isin(REVIEW_STATUSES)].copy() if not result.empty else empty,
        "missing_in_2b": result[result["status"] == "In Register Only"].copy() if not result.empty else empty,
        "missing_in_register": result[result["status"] == "In 2B Only"].copy() if not result.empty else empty,
    }


def read_upload(upload) -> pd.DataFrame:
    suffix = Path(upload.name).suffix.lower()
    content = upload.getvalue()
    if suffix == ".csv":
        return pd.read_csv(BytesIO(content))
    return pd.read_excel(BytesIO(content), sheet_name=0)


# ---------------------------------------------------------------------------
# Excel export
# ---------------------------------------------------------------------------
MONEY_COLUMNS = {
    "taxable", "igst", "cgst", "sgst", "cess", "total", "value_difference_total",
    "Taxable_Value", "IGST", "CGST", "SGST", "Cess", "Invoice_Value",
    "Value_Difference", "Value (Rs.)",
}


def export_results(results: dict[str, pd.DataFrame], progress_callback=None) -> bytes:
    output = BytesIO()
    brand = "9C002C"

    def value_of(frame, column: str = "total") -> float:
        if frame is None or frame.empty or column not in frame.columns:
            return 0.0
        return round(float(frame[column].sum()), 2)

    quality = results.get("quality_checks", pd.DataFrame())
    high_severity = int((quality["Severity"] == "High").sum()) if not quality.empty else 0

    summary = pd.DataFrame([
        {"Metric": "GSTR-2B rows", "Count": len(results["gstr2b_source"]),
         "Value (Rs.)": value_of(results["gstr2b_source"])},
        {"Metric": "Inward register rows", "Count": len(results["inward_register_source"]),
         "Value (Rs.)": value_of(results["inward_register_source"])},
        {"Metric": "Matched", "Count": len(results["matched"]), "Value (Rs.)": value_of(results["matched"])},
        {"Metric": "Review required", "Count": len(results["mismatches"]),
         "Value (Rs.)": value_of(results["mismatches"])},
        {"Metric": "In register only (ITC not in 2B)", "Count": len(results["missing_in_2b"]),
         "Value (Rs.)": value_of(results["missing_in_2b"])},
        {"Metric": "In 2B only (not booked)", "Count": len(results["missing_in_register"]),
         "Value (Rs.)": value_of(results["missing_in_register"])},
        {"Metric": "Suppliers identified", "Count": len(results.get("supplier_directory", pd.DataFrame())),
         "Value (Rs.)": 0.0},
        {"Metric": "High-severity data quality flags", "Count": high_severity, "Value (Rs.)": 0.0},
    ])

    logo_path = Path(__file__).parent / "assets" / "PKP_logo.png"

    sheets = [
        ("Summary", summary),
        ("Supplier_Summary", results.get("supplier_summary")),
        ("Review_Queue", results["mismatches"]),
        ("In_Register_Only", results["missing_in_2b"]),
        ("In_2B_Only", results["missing_in_register"]),
        ("Matched_Details", results.get("matched_details")),
        ("Monthly_Analysis", results.get("monthly_analysis")),
        ("Data_Quality_Checks", quality),
        ("Supplier_Directory", results.get("supplier_directory")),
        ("GSTR_2B_Upload", results["gstr2b_source"]),
        ("Inward_Register_Upload", results["inward_register_source"]),
        ("All_Reconciliation", results["reconciliation"]),
    ]

    with pd.ExcelWriter(output, engine="xlsxwriter") as writer:
        workbook = writer.book
        workbook.set_properties({
            "title": "GST Reconciliation Studio",
            "subject": "GSTR-2B vs Inward Register Reconciliation",
            "author": "PUSHPAK KUMAR",
            "comments": "Professional GST reconciliation and exception reporting workbook.",
        })
        header_format = workbook.add_format({
            "bold": True, "font_color": "#FFFFFF", "bg_color": f"#{brand}",
            "align": "center", "valign": "vcenter", "text_wrap": True,
        })
        review_format = workbook.add_format({"bg_color": "#FFF2CC"})
        high_format = workbook.add_format({"bg_color": "#FCE4E4", "font_color": "#9C0006"})
        money_format = workbook.add_format({"num_format": "#,##0.00"})
        title_format = workbook.add_format({"bold": True, "font_size": 18, "font_color": f"#{brand}"})
        subtitle_format = workbook.add_format({"italic": True, "font_color": "#667085"})
        brand_format = workbook.add_format({"bold": True, "font_color": f"#{brand}", "font_size": 9})
        banner_format = workbook.add_format({"bg_color": "#F9EEF2"})
        link_format = workbook.add_format({"font_color": "#0563C1", "underline": 1})

        for sheet_name, frame in sheets:
            if progress_callback:
                progress_callback(f"Writing {sheet_name} sheet...")
            if frame is None or len(frame.columns) == 0:
                frame = pd.DataFrame({"Note": ["No items in this category."]})
            start_row = 3 if sheet_name == "Summary" else 0
            frame.to_excel(writer, index=False, sheet_name=sheet_name, startrow=start_row)
            worksheet = writer.sheets[sheet_name]
            rows, columns = frame.shape
            header_row = start_row
            worksheet.hide_gridlines(2)
            worksheet.freeze_panes(header_row + 1, 0)
            worksheet.set_tab_color(f"#{brand}")
            worksheet.set_footer("PUSHPAK KUMAR  |  GST Reconciliation Studio")
            worksheet.set_row(header_row, 30)

            for column_index, column_name in enumerate(frame.columns):
                longest = max([len(str(column_name)),
                               *[len(str(value or "")) for value in frame[column_name].head(100)]])
                width = min(max(longest + 2, 12), 34)
                cell_format = money_format if str(column_name) in MONEY_COLUMNS else None
                worksheet.set_column(column_index, column_index, width, cell_format)
                worksheet.write(header_row, column_index, column_name, header_format)

            if "matched_2b_link" in frame.columns and "matched_2b_row_id" in frame.columns:
                link_column = frame.columns.get_loc("matched_2b_link")
                row_column = frame.columns.get_loc("matched_2b_row_id")
                for frame_row, row_id in enumerate(frame.iloc[:, row_column]):
                    if pd.notna(row_id) and str(frame.iloc[frame_row, link_column]).strip():
                        worksheet.write_url(
                            header_row + 1 + frame_row, link_column,
                            f"internal:'GSTR_2B_Upload'!A{int(row_id) + 2}",
                            link_format, string="Open GSTR-2B invoice",
                        )

            if columns and rows:
                worksheet.autofilter(header_row, 0, header_row + rows, columns - 1)

            if "status" in frame.columns and rows:
                status_letter = get_column_letter(frame.columns.get_loc("status") + 1)
                worksheet.conditional_format(
                    header_row + 1, 0, header_row + rows, columns - 1,
                    {"type": "formula",
                     "criteria": f'=${status_letter}{header_row + 2}<>"Matched"',
                     "format": review_format},
                )

            if "Severity" in frame.columns and rows:
                severity_letter = get_column_letter(frame.columns.get_loc("Severity") + 1)
                worksheet.conditional_format(
                    header_row + 1, 0, header_row + rows, columns - 1,
                    {"type": "formula",
                     "criteria": f'=${severity_letter}{header_row + 2}="High"',
                     "format": high_format},
                )

            if sheet_name == "Summary":
                worksheet.merge_range("B1:G1", "GST RECONCILIATION STUDIO", title_format)
                worksheet.merge_range(
                    "B2:G2", "GSTR-2B vs Inward Register | Reconciliation & Exception Report", subtitle_format
                )
                worksheet.merge_range(
                    "B3:G3", "Prepared by PUSHPAK KUMAR  |  GST Reconciliation & Reporting", brand_format
                )
                worksheet.set_row(0, 32)
                worksheet.set_row(1, 22)
                worksheet.set_row(2, 20)
                for row in range(3):
                    worksheet.set_row(row, None, banner_format)
                if logo_path.exists():
                    worksheet.insert_image("A1", str(logo_path), {"x_scale": 0.45, "y_scale": 0.45})

        if progress_callback:
            progress_callback("Excel workbook finalized.")

    return output.getvalue()
